"""Rough local preview of the deck's final state (all shapes visible), for layout QA without LibreOffice"""
import sys
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Ellipse, Rectangle, FancyBboxPatch
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_SHAPE
from pptx.oxml.ns import qn

EMU = 914400


def color_of(fill_or_line, default=None):
    try:
        return "#" + str(fill_or_line.color.rgb)
    except Exception:
        return default


def render(path, out_prefix, only=None):
    prs = Presentation(path)
    for n, slide in enumerate(prs.slides, 1):
        if only and n not in only:
            continue
        fig = plt.figure(figsize=(10, 5.625), dpi=110)
        ax = fig.add_axes([0, 0, 1, 1])
        ax.set_xlim(0, 10)
        ax.set_ylim(5.625, 0)
        ax.set_facecolor("#1C1C1C")
        ax.axis("off")
        fig.patch.set_facecolor("#1C1C1C")
        for sp in slide.shapes:
            x, y, w, h = sp.left / EMU, sp.top / EMU, sp.width / EMU, sp.height / EMU
            if sp.shape_type == MSO_SHAPE_TYPE.PICTURE:
                img = Image.open(sp.image.blob and __import__("io").BytesIO(sp.image.blob))
                ax.imshow(img.convert("RGBA"), extent=(x, x + w, y + h, y), zorder=1)
                continue
            if sp.shape_type == MSO_SHAPE_TYPE.LINE or sp.element.tag == qn("p:cxnSp"):
                flip = sp.element.find(".//" + qn("a:xfrm"))
                fh, fv = flip.get("flipH") == "1", flip.get("flipV") == "1"
                x1, x2 = (x + w, x) if fh else (x, x + w)
                y1, y2 = (y + h, y) if fv else (y, y + h)
                ax.plot([x1, x2], [y1, y2], color=color_of(sp.line, "#BBBBBB"), lw=sp.line.width.pt * 0.8, zorder=2)
                continue
            if sp.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                fill = color_of(sp.fill) if sp.fill.type == 1 else None
                edge = color_of(sp.line) if sp.line.fill.type == 1 else None
                kw = dict(facecolor=fill or "none", edgecolor=edge or "none", lw=(sp.line.width.pt if edge else 0) * 0.8, zorder=2)
                if sp.auto_shape_type == MSO_SHAPE.OVAL:
                    ax.add_patch(Ellipse((x + w / 2, y + h / 2), w, h, angle=sp.rotation, **kw))
                else:
                    ax.add_patch(Rectangle((x, y), w, h, **kw))
            if sp.has_text_frame and sp.text_frame.text.strip():
                tf = sp.text_frame
                runs = [r for p in tf.paragraphs for r in p.runs]
                size = runs[0].font.size.pt if runs and runs[0].font.size else 18
                col = color_of(runs[0].font) if runs else "#FFFFFF"
                align = tf.paragraphs[0].alignment
                ha = {1: "left", 2: "center", 3: "right"}.get(int(align) if align else 1, "left")
                tx = {"left": x, "center": x + w / 2, "right": x + w}[ha]
                ax.text(tx, y + h / 2, tf.text, fontsize=size, color=col, ha=ha,
                        va="center", rotation=-sp.rotation if sp.rotation else 0, zorder=3,
                        fontweight="bold" if runs and runs[0].font.bold else "normal", family=["STIXGeneral"],
                        style="italic" if runs and runs[0].font.italic else "normal")
                ax.add_patch(Rectangle((x, y), w, h, fill=False, ec="#444444", lw=0.3, ls=":", zorder=3))
        fig.savefig(f"{out_prefix}{n:02d}.png")
        plt.close(fig)


if __name__ == "__main__":
    only = {int(a) for a in sys.argv[3:]} or None
    render(sys.argv[1], sys.argv[2], only)
