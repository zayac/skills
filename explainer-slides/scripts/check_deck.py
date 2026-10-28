"""Pre-upload checks for a generated deck: clicks vs [click] marks in the notes, and fonts Google keeps.

Usage: python check_deck.py deck.pptx
Exit code 1 if anything needs fixing.
"""
import sys
from pptx import Presentation
from pptx.oxml.ns import qn

# Fonts confirmed to survive .pptx -> Google Slides conversion. Anything else may silently become
# Arial: Google keeps a font only if its editors offer it (Noto Serif is not offered, only its
# script variants such as Noto Serif Thai, so it fell back to Arial).
CONFIRMED = {"Arial", "STIX Two Text", "STIX Two Math"}
# Seen in the Slides font menu but not yet put through a conversion: probably fine, check after.
IN_MENU = {
    "Roboto", "Roboto Slab", "Lato", "Montserrat", "Merriweather", "Lora", "Source Serif 4", "Noto Sans",
    "Nunito", "Oswald", "Lexend", "EB Garamond", "JetBrains Mono", "Libertinus Math", "Calibri", "Cambria",
    "Cambria Math", "Consolas", "Courier New", "Comfortaa",
}


def main(path):
    prs = Presentation(path)
    problems = 0
    fonts = set()
    for n, slide in enumerate(prs.slides, 1):
        clicks = slide._element.xml.count('delay="indefinite"')
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        marks = notes.count("[click]")
        if clicks != marks:
            print(f"slide {n}: {clicks} animation clicks but {marks} [click] marks in the notes")
            problems += 1
        if not notes.strip():
            print(f"slide {n}: no speaker notes")
        for latin in slide._element.iter(qn("a:latin")):
            fonts.add(latin.get("typeface"))
    for font in sorted(fonts):
        if font in CONFIRMED:
            continue
        if font in IN_MENU:
            print(f"font {font!r}: in Google's menu but conversion unverified; check it in the converted deck")
        else:
            print(f"font {font!r}: not known to survive conversion; Google may replace it with Arial")
            problems += 1
    print("OK" if not problems else f"{problems} problem(s)")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1]))
